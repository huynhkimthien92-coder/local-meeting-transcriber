"""
Xuất biên bản ra file DOCX theo template chuẩn — dùng python-docx (free,
open-source, không cần Microsoft Word cài trên máy).
"""
from __future__ import annotations

import re
from pathlib import Path

from ..pipeline.models import MeetingJob


def export_to_docx(job: MeetingJob, output_path: str | Path) -> Path:
    try:
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as e:
        raise RuntimeError("Chưa cài đặt python-docx. Chạy: pip install python-docx") from e

    output_path = Path(output_path)
    doc = Document()

    # --- Tiêu đề ---
    title = doc.add_heading(f"BIÊN BẢN CUỘC HỌP", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(job.title)
    run.bold = True
    run.font.size = Pt(14)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Ngày tạo: {job.created_at}").italic = True

    doc.add_paragraph()  # dòng trống

    # --- Nội dung biên bản (parse markdown đơn giản từ summary_markdown) ---
    if job.summary_markdown:
        _append_markdown(doc, job.summary_markdown)
    else:
        doc.add_paragraph("(Chưa có bản tóm tắt)")

    # --- Phụ lục: transcript đầy đủ có tên người nói ---
    doc.add_page_break()
    doc.add_heading("Phụ lục: Bản ghi chi tiết", level=1)
    for seg in job.segments:
        speaker_display = job.speaker_names.get(seg.speaker, seg.speaker)
        p = doc.add_paragraph()
        ts = _format_timestamp(seg.start)
        p.add_run(f"[{ts}] ").italic = True
        p.add_run(f"{speaker_display}: ").bold = True
        p.add_run(seg.text)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def _format_timestamp(seconds: float) -> str:
    m, s = divmod(int(seconds), 60)
    h, m = divmod(m, 60)
    if h:
        return f"{h:02d}:{m:02d}:{s:02d}"
    return f"{m:02d}:{s:02d}"


_HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)")
_TABLE_ROW_RE = re.compile(r"^\|(.+)\|$")
# Model đôi khi không theo đúng "## Tiêu đề" như prompt yêu cầu (đặc biệt các
# model nhỏ như llama3.2:1b), mà tự ý dùng "**Tiêu đề:**" — nhận diện thêm
# dạng này để không bị lộ dấu ** thừa trong file Word.
_BOLD_ONLY_LINE_RE = re.compile(r"^\*\*(.+?)\*\*:?\s*$")
_INLINE_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _append_markdown(doc, markdown_text: str) -> None:
    """Parser markdown tối giản: đủ cho định dạng biên bản do LLM sinh ra
    (heading ##, bullet -, bảng | | |, và **bold** cả dạng nguyên dòng lẫn
    chen giữa câu). Không cần thư viện markdown đầy đủ vì output của prompt
    trong summarize.py đã cố định cấu trúc này — đây chỉ là lớp đệm cho lúc
    model (nhất là model nhỏ) không bám sát 100% format đã yêu cầu."""
    lines = markdown_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        heading_match = _HEADING_RE.match(line)
        if heading_match:
            doc.add_heading(heading_match.group(2), level=min(len(heading_match.group(1)) + 1, 4))
            i += 1
            continue

        bold_only_match = _BOLD_ONLY_LINE_RE.match(line)
        if bold_only_match:
            doc.add_heading(bold_only_match.group(1).strip(), level=2)
            i += 1
            continue

        if line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_inline_bold(p, line[2:])
            i += 1
            continue

        if _TABLE_ROW_RE.match(line):
            table_lines = []
            while i < len(lines) and _TABLE_ROW_RE.match(lines[i].strip()):
                table_lines.append(lines[i].strip())
                i += 1
            _append_table(doc, table_lines)
            continue

        p = doc.add_paragraph()
        _add_runs_with_inline_bold(p, line)
        i += 1


def _add_runs_with_inline_bold(paragraph, text: str) -> None:
    """Thêm text vào paragraph, biến các đoạn **in đậm** thành run in đậm
    thật trong Word thay vì để lộ dấu ** thô."""
    pos = 0
    for m in _INLINE_BOLD_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _append_table(doc, table_lines: list[str]) -> None:
    rows = []
    for line in table_lines:
        cells = [c.strip() for c in line.strip("|").split("|")]
        # bỏ dòng phân cách kiểu |---|---|
        if all(set(c) <= {"-", ":"} for c in cells if c):
            continue
        rows.append(cells)
    if not rows:
        return
    n_cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Light Grid Accent 1"
    for r, row_cells in enumerate(rows):
        for c, cell_text in enumerate(row_cells):
            if c < n_cols:
                table.cell(r, c).text = cell_text
